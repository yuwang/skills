#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1,<4"]
# ///

"""Convert tabular data between JSON, JSONL, CSV, and XLSX."""

from __future__ import annotations

import argparse
import csv
import json
import os
import sys
import tempfile
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Callable, Iterable
from zipfile import BadZipFile


SUPPORTED_FORMATS = {".json": "json", ".jsonl": "jsonl", ".csv": "csv", ".xlsx": "xlsx"}
SPREADSHEET_FORMULA_PREFIXES = ("=", "+", "-", "@")


class ConversionError(Exception):
    """An input or conversion error that should be shown without a traceback."""


def detect_format(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        return SUPPORTED_FORMATS[suffix]
    except KeyError as exc:
        supported = ", ".join(sorted(SUPPORTED_FORMATS))
        raise ConversionError(f"不支持扩展名 {suffix or '<无>'}；支持：{supported}") from exc


def validate_headers(headers: Iterable[Any]) -> list[str]:
    normalized: list[str] = []
    seen: set[str] = set()
    for index, header in enumerate(headers, start=1):
        if header is None or str(header).strip() == "":
            raise ConversionError(f"第 {index} 列的表头为空")
        name = str(header)
        if name in seen:
            raise ConversionError(f"表头重复：{name}")
        if any(part == "" for part in name.split(".")):
            raise ConversionError(f"表头包含无效的点号路径：{name}")
        seen.add(name)
        normalized.append(name)
    validate_header_paths(normalized)
    return normalized


def validate_header_paths(headers: Iterable[str]) -> None:
    header_set = set(headers)
    for header in header_set:
        parts = header.split(".")
        for length in range(1, len(parts)):
            prefix = ".".join(parts[:length])
            if prefix in header_set:
                raise ConversionError(f"字段路径冲突：{prefix} 与 {header}")


def flatten_object(value: dict[str, Any], prefix: str = "") -> dict[str, Any]:
    flattened: dict[str, Any] = {}
    if not value and prefix:
        return {prefix: {}}

    for raw_key, item in value.items():
        key = str(raw_key)
        if not key or "." in key:
            path = f"{prefix}.{key}" if prefix else key
            raise ConversionError(f"JSON 字段名不能留空或包含点号：{path or '<空>'}")
        path = f"{prefix}.{key}" if prefix else key
        nested = flatten_object(item, path) if isinstance(item, dict) else {path: item}
        for nested_path, nested_value in nested.items():
            if nested_path in flattened:
                raise ConversionError(f"字段路径冲突：{nested_path}")
            flattened[nested_path] = nested_value
    return flattened


def decode_complex_cell(value: Any) -> Any:
    if value is None or value == "":
        return None
    if not isinstance(value, str):
        return value

    stripped = value.strip()
    if not (
        (stripped.startswith("[") and stripped.endswith("]"))
        or (stripped.startswith("{") and stripped.endswith("}"))
    ):
        return value
    try:
        decoded = json.loads(stripped)
    except json.JSONDecodeError:
        return value
    return decoded if isinstance(decoded, (list, dict)) else value


def unflatten_record(record: dict[str, Any]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for header, raw_value in record.items():
        parts = header.split(".")
        cursor = result
        for index, part in enumerate(parts[:-1], start=1):
            if part not in cursor:
                cursor[part] = {}
            elif not isinstance(cursor[part], dict):
                prefix = ".".join(parts[:index])
                raise ConversionError(f"字段路径冲突：{prefix} 与 {header}")
            cursor = cursor[part]

        leaf = parts[-1]
        if leaf in cursor:
            raise ConversionError(f"字段路径冲突：{header}")
        cursor[leaf] = decode_complex_cell(raw_value)
    return result


def ensure_object_records(value: Any, source: str) -> list[dict[str, Any]]:
    if isinstance(value, dict):
        return [value]
    if not isinstance(value, list):
        raise ConversionError(f"{source} 顶层必须是对象或对象数组")
    for index, item in enumerate(value, start=1):
        if not isinstance(item, dict):
            raise ConversionError(f"{source} 第 {index} 条记录不是对象")
    return value


def read_json(path: Path, _sheet: str | None) -> list[dict[str, Any]]:
    try:
        with path.open(encoding="utf-8-sig") as handle:
            value = json.load(handle)
    except json.JSONDecodeError as exc:
        raise ConversionError(
            f"JSON 解析失败：第 {exc.lineno} 行第 {exc.colno} 列，{exc.msg}"
        ) from exc
    return ensure_object_records(value, "JSON")


def read_jsonl(path: Path, _sheet: str | None) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    with path.open(encoding="utf-8-sig") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            try:
                value = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ConversionError(f"JSONL 第 {line_number} 行解析失败：{exc.msg}") from exc
            if not isinstance(value, dict):
                raise ConversionError(f"JSONL 第 {line_number} 行不是对象")
            records.append(value)
    return records


def read_csv(path: Path, _sheet: str | None) -> list[dict[str, Any]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        try:
            raw_headers = next(reader)
        except StopIteration as exc:
            raise ConversionError("CSV 为空，缺少表头") from exc
        headers = validate_headers(raw_headers)
        records: list[dict[str, Any]] = []
        for line_number, values in enumerate(reader, start=2):
            if len(values) != len(headers):
                raise ConversionError(
                    f"CSV 第 {line_number} 行有 {len(values)} 列，表头有 {len(headers)} 列"
                )
            records.append(unflatten_record(dict(zip(headers, values))))
    return records


def read_xlsx(path: Path, sheet: str | None) -> list[dict[str, Any]]:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException) as exc:
        raise ConversionError(f"XLSX 文件无效：{path}（{exc}）") from exc
    try:
        if sheet is None:
            worksheet = workbook.active
        elif sheet in workbook.sheetnames:
            worksheet = workbook[sheet]
        else:
            available = ", ".join(workbook.sheetnames)
            raise ConversionError(f"工作表不存在：{sheet}；可用工作表：{available}")

        rows = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise ConversionError(f"工作表为空：{worksheet.title}") from exc
        headers = validate_headers(raw_headers)

        records: list[dict[str, Any]] = []
        for row_number, raw_values in enumerate(rows, start=2):
            values = list(raw_values)
            if len(values) > len(headers) and any(
                value is not None for value in values[len(headers) :]
            ):
                raise ConversionError(f"工作表第 {row_number} 行包含表头范围外的数据")
            values = (values + [None] * len(headers))[: len(headers)]
            if all(value is None for value in values):
                continue
            records.append(unflatten_record(dict(zip(headers, values))))
        return records
    finally:
        workbook.close()


READERS: dict[str, Callable[[Path, str | None], list[dict[str, Any]]]] = {
    "json": read_json,
    "jsonl": read_jsonl,
    "csv": read_csv,
    "xlsx": read_xlsx,
}


def json_safe(value: Any) -> Any:
    if isinstance(value, dict):
        return {key: json_safe(item) for key, item in value.items()}
    if isinstance(value, list):
        return [json_safe(item) for item in value]
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return value


def table_cell(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        value = json.dumps(json_safe(value), ensure_ascii=False, separators=(",", ":"))
    elif isinstance(value, (date, datetime, time)):
        value = value.isoformat()
    if isinstance(value, str) and value.startswith(SPREADSHEET_FORMULA_PREFIXES):
        return "'" + value
    return value


def flatten_records(records: list[dict[str, Any]]) -> tuple[list[str], list[dict[str, Any]]]:
    rows: list[dict[str, Any]] = []
    headers: list[str] = []
    seen: set[str] = set()
    for record in records:
        flattened = flatten_object(record)
        rows.append(flattened)
        for header in flattened:
            if header not in seen:
                headers.append(header)
                seen.add(header)
    if not headers:
        raise ConversionError("没有可写入表格的字段")
    validate_header_paths(headers)
    return headers, rows


def write_json(records: list[dict[str, Any]], path: Path, _sheet: str | None) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        json.dump(json_safe(records), handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def write_jsonl(records: list[dict[str, Any]], path: Path, _sheet: str | None) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        for record in records:
            json.dump(
                json_safe(record),
                handle,
                ensure_ascii=False,
                separators=(",", ":"),
            )
            handle.write("\n")


def write_csv(records: list[dict[str, Any]], path: Path, _sheet: str | None) -> None:
    headers, rows = flatten_records(records)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow([table_cell(header) for header in headers])
        for row in rows:
            writer.writerow([table_cell(row.get(header)) for header in headers])


def validate_sheet_name(sheet: str) -> None:
    invalid = set('[]:*?/\\')
    if not sheet or len(sheet) > 31 or any(character in invalid for character in sheet):
        raise ConversionError(f"无效的 XLSX 工作表名称：{sheet!r}")


def write_xlsx(records: list[dict[str, Any]], path: Path, sheet: str | None) -> None:
    from openpyxl import Workbook

    headers, rows = flatten_records(records)
    sheet_name = sheet or "Data"
    validate_sheet_name(sheet_name)
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet(title=sheet_name)
    worksheet.append([table_cell(header) for header in headers])
    for row in rows:
        worksheet.append([table_cell(row.get(header)) for header in headers])
    workbook.save(path)


WRITERS: dict[str, Callable[[list[dict[str, Any]], Path, str | None], None]] = {
    "json": write_json,
    "jsonl": write_jsonl,
    "csv": write_csv,
    "xlsx": write_xlsx,
}


def atomic_write(
    records: list[dict[str, Any]], output: Path, output_format: str, sheet: str | None
) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    file_descriptor, temporary_name = tempfile.mkstemp(
        dir=output.parent,
        prefix=f".{output.stem}.",
        suffix=output.suffix,
    )
    os.close(file_descriptor)
    temporary = Path(temporary_name)
    try:
        WRITERS[output_format](records, temporary, sheet)
        os.replace(temporary, output)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise


def convert(input_path: Path, output_path: Path, sheet: str | None, force: bool) -> None:
    if not input_path.is_file():
        raise ConversionError(f"输入文件不存在：{input_path}")
    input_format = detect_format(input_path)
    output_format = detect_format(output_path)
    if input_format == output_format:
        raise ConversionError(f"输入和输出格式相同：{input_format}")
    if input_path.resolve() == output_path.resolve():
        raise ConversionError("输入和输出不能是同一个文件")
    if output_path.exists() and not force:
        raise ConversionError(f"输出文件已存在；使用 --force 覆盖：{output_path}")
    if sheet is not None and "xlsx" not in (input_format, output_format):
        raise ConversionError("--sheet 仅用于 XLSX 输入或输出")

    records = READERS[input_format](input_path, sheet if input_format == "xlsx" else None)
    atomic_write(records, output_path, output_format, sheet if output_format == "xlsx" else None)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="在 JSON、JSONL、CSV、XLSX 之间转换表格数据。"
    )
    parser.add_argument("input", type=Path, help="输入文件")
    parser.add_argument("output", type=Path, help="输出文件")
    parser.add_argument("--sheet", help="读取或创建的 XLSX 工作表名称")
    parser.add_argument("--force", action="store_true", help="覆盖已存在的输出文件")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    try:
        convert(args.input, args.output, args.sheet, args.force)
    except (ConversionError, OSError, ValueError) as exc:
        print(f"转换失败：{exc}", file=sys.stderr)
        return 1
    print(f"转换完成：{args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
